import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout_box(doc, text, title="IMPORTANT COMPLIANCE / NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(27, 54, 93)
    
    run_text = p.add_run(text)
    run_text.font.name = "Segoe UI"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(51, 51, 51)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table(table, header_bg="1B365D", alt_bg="F9FBFD"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header styling
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.name = "Segoe UI"
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(10)
    
    # Body rows
    for i, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if i % 2 == 0 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Segoe UI"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(40, 40, 40)

def generate_documentation():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Bank_Guide_AI — Technical & System Documentation")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential — Internal Banking AI Architecture & SOP Assistant")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(128, 128, 128)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    
    # -------------------------------------------------------------
    # COVER / TITLE BLOCK
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(4)
    run_main_title = title_p.add_run("Bank_Guide_AI")
    run_main_title.bold = True
    run_main_title.font.name = "Segoe UI"
    run_main_title.font.size = Pt(28)
    run_main_title.font.color.rgb = RGBColor(27, 54, 93) # Navy Blue
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(18)
    run_sub = subtitle_p.add_run("Bilingual Banking Standard Operating Procedures (SOP) Knowledge Assistant\nwith GPU-Accelerated Hybrid RAG (Dense BGE-M3 + Sparse BM25 + RRF)")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(74, 119, 122)
    
    # Metadata Table
    meta_tbl = doc.add_table(rows=6, cols=2)
    meta_data = [
        ("Parameter", "System Specification"),
        ("Document Version", "2.0 (Production Release — Hybrid Retrieval Enabled)"),
        ("Classification", "Internal Enterprise Banking Architecture Specification"),
        ("Author & Engineering Lead", "Ahmed Esam (AI Engineering & Architecture)"),
        ("Knowledge Domain", "Standard Operating Procedures (SOP), Compliance & Internal Manuals"),
        ("Supported Languages", "Arabic (Native Source Documents) & English (Cross-Lingual Queries)")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        meta_tbl.cell(row_idx, 0).text = k
        meta_tbl.cell(row_idx, 1).text = v
    style_table(meta_tbl, header_bg="2B4C7E", alt_bg="F4F7FB")
    meta_tbl.columns[0].width = Inches(2.2)
    meta_tbl.columns[1].width = Inches(4.3)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary & Business Objectives", level=1)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True
        
    p = doc.add_paragraph(
        "Commercial and retail banking operations depend heavily on rigorous compliance with internal "
        "Standard Operating Procedures (SOPs), circulars, and departmental manuals. However, these documents "
        "often span hundreds of dense, bureaucratic Arabic pages, scanned PDF circulars, and intricate "
        "multi-step workflows. Branch officers, compliance auditors, and warehouse managers face significant "
        "delays when manually locating specific procedural clauses—such as daily postal cut-off times, lost keycard "
        "penalties, CCTV extraction permissions, and fixed asset disposal protocols."
    )
    p.paragraph_format.space_after = Pt(8)
    
    p = doc.add_paragraph(
        "Bank_Guide_AI is an enterprise-grade Retrieval-Augmented Generation (RAG) system engineered to solve this "
        "challenge. It delivers instant, zero-hallucination, bilingual procedural guidance by combining:"
    )
    p.paragraph_format.space_after = Pt(4)
    
    bullets = [
        ("⚡ GPU-Accelerated Dense Embeddings: ", "Utilizes BAAI/bge-m3 on NVIDIA CUDA for rapid semantic vector representation across 100+ languages."),
        ("🔍 Sparse BM25 Keyword Search: ", "Provides exact lexical matching for banking acronyms (BPM, FIFO, ATM), article numbers, and numerical policies."),
        ("🔀 Hybrid Reciprocal Rank Fusion (RRF): ", "Harmonizes dense semantic nuance with exact keyword precision with user-tunable weighting."),
        ("📑 OCR-Aware Fallback Pipeline: ", "Automatically parses native PDF text layers via PyMuPDF while invoking Tesseract OCR (ara+eng) for scanned circulars and rasterized tables."),
        ("🛡️ Grounded Citation Governance: ", "Every LLM response is strictly constrained to retrieved context, citing the exact document name, page number, and extraction method.")
    ]
    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(b_title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(27, 54, 93)
        r2 = bp.add_run(b_desc)
        
    add_callout_box(
        doc,
        "The system enforces a strict anti-hallucination policy: If a requested policy or procedure is not present in "
        "the ingested knowledge base, the assistant explicitly states that the manuals do not cover the subject, "
        "preventing operational or compliance risks.",
        title="GOVERNANCE & COMPLIANCE STANDARD"
    )
    
    # -------------------------------------------------------------
    # 2. INGESTED KNOWLEDGE BASE
    # -------------------------------------------------------------
    h1 = doc.add_heading("2. Ingested Procedural Knowledge Base", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True

    p = doc.add_paragraph("The system currently indexes three core banking procedure manuals located in the data repository:")
    p.paragraph_format.space_after = Pt(6)
    
    kb_tbl = doc.add_table(rows=4, cols=3)
    kb_data = [
        ("Manual Name", "Arabic Title & Scope", "Key Core Topics & Operational Clauses"),
        (
            "Central Mail & Files Manual\n(Central_Mail_and_Files_Unit_Procedures_Manual.pdf)",
            "دليل إجراءات وحدة البريد المركزي والملفات\n(Administration / Operations)",
            "• Daily BPM submission cut-off at 1:30 PM\n• Prohibition of sending debit card + PIN in same parcel\n• Archiving of approved credit files (Doxis / physical)\n• Domestic & international courier contracts"
        ),
        (
            "Central Alarm Tasks Manual\n(Central_Alarm_Tasks_and_Procedures_Manual.pdf)",
            "دليل إجراءات وحدة الإنذار المركزي\n(Security & Branch Safety)",
            "• 10 JOD replacement penalty for lost access cards\n• Exemption criteria (technical defect or after 2 years)\n• CCTV footage extraction protocols & security approvals\n• Fire and anti-burglary alarm monitoring cycles"
        ),
        (
            "Assets & Warehouse Operations Manual\n(Assets_and_Warehouse_Operations_Manual.pdf)",
            "دليل إجراءات وحدة الموجودات وعمليات المستودعات\n(Supply Chain & Logistics)",
            "• FIFO (First-In, First-Out) inventory dispatch policy\n• Annual & surprise physical inventory committees\n• Scrap, damaged asset disposal, and donation approvals\n• Goods Receipt Note (GRN) documentation standards"
        )
    ]
    for row_idx, row in enumerate(kb_data):
        for col_idx, text in enumerate(row):
            kb_tbl.cell(row_idx, col_idx).text = text
    style_table(kb_tbl, header_bg="1B365D", alt_bg="F9FBFD")
    kb_tbl.columns[0].width = Inches(2.1)
    kb_tbl.columns[1].width = Inches(2.1)
    kb_tbl.columns[2].width = Inches(2.3)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 3. SYSTEM ARCHITECTURE & DATA FLOW
    # -------------------------------------------------------------
    h1 = doc.add_heading("3. End-to-End System Architecture", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True

    p = doc.add_paragraph(
        "The architecture is engineered into three modular, high-performance pipelines: "
        "(1) Ingestion & Vectorization, (2) Hybrid Retrieval Engine, and (3) Grounded Generation."
    )
    p.paragraph_format.space_after = Pt(8)

    arch_tbl = doc.add_table(rows=4, cols=3)
    arch_data = [
        ("Pipeline Stage", "Core Technologies", "Responsibilities & Key Mechanisms"),
        (
            "1. Ingestion &\nPreprocessing",
            "• PyMuPDF (fitz)\n• Tesseract OCR (ara+eng)\n• 5 Chunking Strategies\n• HuggingFace BGE-M3 (CUDA)",
            "• Native text extraction with automatic fallback to OCR\n• JSON extraction cache to prevent redundant OCR\n• 5 Chunking algorithms (recursive, arabic_paragraph, etc.)\n• GPU-accelerated batch embedding into 1024-dim vectors\n• Persistent storage in ChromaDB vector store"
        ),
        (
            "2. Hybrid Retrieval\n& Fusion Engine",
            "• ChromaDB (Dense Vector)\n• BM25Retriever (rank-bm25)\n• Custom HybridEnsembleRetriever\n• Reciprocal Rank Fusion (RRF)",
            "• Dense vector similarity / MMR search for semantic intent\n• BM25 sparse keyword search for acronyms/article numbers\n• Weighted Reciprocal Rank Fusion (RRF formula with constant c=60)\n• Dynamic UI weight ratio slider (0.0 to 1.0)\n• Document source filtering and top-k tuning"
        ),
        (
            "3. Grounded Generation\n& User Interface",
            "• Groq Cloud API\n• Qwen 2.5 / ALLaM / Llama-3.3\n• LangChain Core Prompts\n• Streamlit Interactive App",
            "• Dynamic prompt synthesis with system anti-hallucination guardrails\n• Cross-lingual comprehension (Arabic context -> English/Arabic response)\n• Source attribution expander with document name, page, and preview\n• Full multi-turn conversation memory within session state"
        )
    ]
    for row_idx, row in enumerate(arch_data):
        for col_idx, text in enumerate(row):
            arch_tbl.cell(row_idx, col_idx).text = text
    style_table(arch_tbl, header_bg="2B4C7E", alt_bg="F4F7FB")
    arch_tbl.columns[0].width = Inches(1.8)
    arch_tbl.columns[1].width = Inches(2.2)
    arch_tbl.columns[2].width = Inches(2.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 4. DEEP DIVE: HYBRID RETRIEVAL & RECIPROCAL RANK FUSION (RRF)
    # -------------------------------------------------------------
    h1 = doc.add_heading("4. Hybrid Retrieval & Reciprocal Rank Fusion (RRF)", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True

    p = doc.add_paragraph(
        "A critical innovation in Bank_Guide_AI v2.0 is the implementation of Hybrid Retrieval. Pure semantic "
        "embeddings can sometimes miss exact numerical article codes or banking abbreviations (such as BPM or FIFO), "
        "while pure keyword search fails on semantic paraphrasing and cross-lingual translation. Hybrid retrieval "
        "harmonizes both paradigms using Reciprocal Rank Fusion (RRF)."
    )
    p.paragraph_format.space_after = Pt(8)

    # RRF Formula Box
    add_callout_box(
        doc,
        "RRF Score Formulation:\n"
        "Score(d) = w_dense * (1 / (c + rank_dense(d) + 1)) + w_sparse * (1 / (c + rank_sparse(d) + 1))\n\n"
        "Where:\n"
        "• rank(d) is the 0-indexed position of document chunk d in the respective retriever's ranked list.\n"
        "• c = 60 is the standard rank smoothing constant to prevent outliers from dominating the score.\n"
        "• w_dense and w_sparse are user-tunable weights (default: 0.5 dense, 0.5 sparse) configured via the UI slider.",
        title="MATHEMATICAL FORMULATION — RECIPROCAL RANK FUSION (RRF)"
    )

    h2 = doc.add_heading("4.1 Retrieval Modes Comparison", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    for r in h2.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(74, 119, 122)
        r.font.size = Pt(12)
        r.bold = True

    mode_tbl = doc.add_table(rows=4, cols=4)
    mode_data = [
        ("Retrieval Mode", "Underlying Technology", "Ideal Use Cases", "Strengths / Limitations"),
        (
            "🔀 Hybrid (Default)",
            "ChromaDB (BGE-M3) +\nBM25 (rank-bm25) via RRF",
            "General banking inquiries, mixed Arabic/English queries, policy lookups with specific terms.",
            "Combines conceptual semantic understanding with exact lexical matching; highest empirical accuracy."
        ),
        (
            "🧠 Semantic Vector",
            "ChromaDB + BGE-M3\n(MMR or Cosine Similarity)",
            "Conceptual questions, descriptive workflows, paraphrased cross-lingual questions.",
            "Handles synonyms and cross-language translation flawlessly; may rank exact code matches lower."
        ),
        (
            "🔍 Keyword (BM25)",
            "BM25 Sparse Lexical\nInverted Index",
            "Exact form IDs, penalty numbers ('10 دنانير'), section numbers, specific system names (Doxis, BPM).",
            "Fast, exact keyword precision; unable to perform cross-lingual semantic bridging or synonym matching."
        )
    ]
    for row_idx, row in enumerate(mode_data):
        for col_idx, text in enumerate(row):
            mode_tbl.cell(row_idx, col_idx).text = text
    style_table(mode_tbl, header_bg="1B365D", alt_bg="F9FBFD")
    mode_tbl.columns[0].width = Inches(1.5)
    mode_tbl.columns[1].width = Inches(1.6)
    mode_tbl.columns[2].width = Inches(1.8)
    mode_tbl.columns[3].width = Inches(1.6)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 5. CHUNKING STRATEGIES & OCR INGESTION
    # -------------------------------------------------------------
    h1 = doc.add_heading("5. Chunking Strategies & Ingestion Engineering", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True

    p = doc.add_paragraph(
        "Banking documents possess unique structural characteristics: numbered procedural steps, multi-level Arabic "
        "sub-headings, and formal administrative tables. To support different retrieval needs, Bank_Guide_AI implements "
        "five interchangeable chunking strategies:"
    )
    p.paragraph_format.space_after = Pt(6)

    chk_tbl = doc.add_table(rows=6, cols=3)
    chk_data = [
        ("Strategy Key", "Implementation Mechanism", "Best Suited For"),
        ("recursive_character\n(Default)", "Splits hierarchically on paragraphs (\\n\\n), lines (\\n), sentences (.), and words.", "General SOP content; balanced chunk sizes with consistent context boundaries."),
        ("arabic_paragraph", "Regex-aware split on Arabic numerals (-1, 1-, أولاً, ثانياً) and double newlines.", "Procedural step-by-step manuals where individual clauses must stay atomic."),
        ("markdown_heading", "Splits along structural markdown/document headers (#, ##, المادة, البند).", "Hierarchically organized manuals and policy chapters."),
        ("character", "Strict sliding-window slicing with exact character length and overlap.", "Dense, unstructured narrative texts where structural delimiters are absent."),
        ("token_based", "Token-budgeted splitting using tiktoken (cl100k_base tokenizer).", "Strict LLM context window budgeting and precise token control.")
    ]
    for row_idx, row in enumerate(chk_data):
        for col_idx, text in enumerate(row):
            chk_tbl.cell(row_idx, col_idx).text = text
    style_table(chk_tbl, header_bg="2B4C7E", alt_bg="F4F7FB")
    chk_tbl.columns[0].width = Inches(1.8)
    chk_tbl.columns[1].width = Inches(2.4)
    chk_tbl.columns[2].width = Inches(2.3)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 6. VERIFICATION TEST MATRIX & BENCHMARK QUERIES
    # -------------------------------------------------------------
    h1 = doc.add_heading("6. Verification Test Matrix & Benchmark Queries", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True

    p = doc.add_paragraph("The following test benchmark validates bilingual retrieval accuracy, cross-lingual translation, and exact citation attribution:")
    p.paragraph_format.space_after = Pt(6)

    test_tbl = doc.add_table(rows=5, cols=4)
    test_data = [
        ("Language", "Test Prompt / Inquiry", "Expected Grounded Fact", "Grounded Source & Page"),
        (
            "English",
            "What is the daily cut-off time for submitting outgoing mail requests through BPM?",
            "1:30 PM (Daily cut-off time for BPM workflow submission)",
            "Central Mail Manual\n(Page 5, Extraction: fitz)"
        ),
        (
            "Arabic",
            "ما هي الغرامة المالية في حال فقدان بطاقة الدخول؟ ومتى يُعفى الموظف منها؟",
            "10 دنانير أردنية؛ يُعفى الموظف في حال الخلل الفني أو بعد مرور سنتين من الإصدار",
            "Central Alarm Manual\n(Page 5, Extraction: fitz)"
        ),
        (
            "English",
            "What inventory dispatch method is used in the bank's warehouses?",
            "FIFO (First-In, First-Out) method for outgoing supply management",
            "Assets & Warehouse Manual\n(Page 5, Extraction: fitz)"
        ),
        (
            "Arabic",
            "هل يجوز إرسال البطاقة المصرفية والرقم السري في شحنة بريدية واحدة للخارج؟",
            "لا يجوز الجمع بينهما إطلاقاً في شحنة واحدة لأسباب أمنية مصرفية",
            "Central Mail Manual\n(Page 9, Extraction: fitz)"
        )
    ]
    for row_idx, row in enumerate(test_data):
        for col_idx, text in enumerate(row):
            test_tbl.cell(row_idx, col_idx).text = text
    style_table(test_tbl, header_bg="1B365D", alt_bg="F9FBFD")
    test_tbl.columns[0].width = Inches(0.9)
    test_tbl.columns[1].width = Inches(2.2)
    test_tbl.columns[2].width = Inches(1.8)
    test_tbl.columns[3].width = Inches(1.6)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 7. CODEBASE REFERENCE & MODULE INVENTORY
    # -------------------------------------------------------------
    h1 = doc.add_heading("7. Codebase Structure & Module Inventory", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True

    code_tbl = doc.add_table(rows=11, cols=2)
    code_data = [
        ("Module Filepath", "Technical Responsibility"),
        ("config.py", "Global parameters: model names, paths, retrieval modes, default weights, CUDA device flags."),
        ("app.py", "Streamlit UI with sidebar ingestion, model selection, retrieval mode controls, and chat history."),
        ("requirements.txt", "Project dependencies: langchain, chromadb, sentence-transformers, rank-bm25, streamlit."),
        ("ingestion/ocr_loader.py", "PyMuPDF text extraction with Tesseract OCR fallback and JSON extraction caching."),
        ("ingestion/chunking.py", "Implementation of all 5 chunking algorithms with metadata preservation."),
        ("ingestion/embeddings.py", "BAAI/bge-m3 embedding model initialization with GPU/CUDA acceleration and batching."),
        ("ingestion/ingest.py", "Orchestrator and CLI runner for batch document ingestion into ChromaDB."),
        ("retrieval/vectorstore.py", "Chroma vectorstore initialization, connection management, and batch document upsert."),
        ("retrieval/retriever.py", "HybridEnsembleRetriever (RRF), BM25 keyword retriever, and Chroma semantic retriever."),
        ("generation/generator.py", "Groq LLM invocation, system prompt formatting, context truncation, and answer synthesis.")
    ]
    for row_idx, row in enumerate(code_data):
        code_tbl.cell(row_idx, 0).text = row[0]
        code_tbl.cell(row_idx, 1).text = row[1]
    style_table(code_tbl, header_bg="2B4C7E", alt_bg="F4F7FB")
    code_tbl.columns[0].width = Inches(2.2)
    code_tbl.columns[1].width = Inches(4.3)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 8. INSTALLATION, CONFIGURATION & RUNTIME GUIDE
    # -------------------------------------------------------------
    h1 = doc.add_heading("8. Installation, Configuration & Runtime Guide", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(27, 54, 93)
        r.font.size = Pt(16)
        r.bold = True

    p = doc.add_paragraph("Follow these operational steps to deploy and run the system:")
    p.paragraph_format.space_after = Pt(4)

    steps = [
        ("1. Environment Setup: ", "Create a Python 3.10 virtual environment and install dependencies via `pip install -r requirements.txt`."),
        ("2. API Key Configuration: ", "Copy `.env.example` to `.env` and configure `GROQ_API_KEY=gsk_...` with your Groq API credentials."),
        ("3. Knowledge Base Ingestion: ", "Execute `python -m ingestion.ingest --strategy recursive_character` or trigger ingestion directly in the Streamlit sidebar."),
        ("4. Application Launch: ", "Launch the web interface via `streamlit run app.py` and access http://localhost:8501 in your browser.")
    ]
    for st_title, st_desc in steps:
        sp = doc.add_paragraph(style='List Bullet')
        sp.paragraph_format.space_after = Pt(3)
        r1 = sp.add_run(st_title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(27, 54, 93)
        r2 = sp.add_run(st_desc)

    # Save to file in project root
    project_root = Path(__file__).resolve().parent
    output_path = project_root / "Bank_Guide_AI_Documentation.docx"
    doc.save(str(output_path))
    print(f"Documentation saved successfully to: {output_path}")

if __name__ == "__main__":
    generate_documentation()
